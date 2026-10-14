"""
extract_student_ledger.py
=========================
Extracts all 7 columns from Student Ledger PDFs:
  Date | Code | Description | Program | Payment Period | Notes | Amount

Run:
    python extract_student_ledger.py <source_folder> <output.xlsx>

Or import and call extract_ledger_folder() directly.
"""

import re
import sys
import os
from pathlib import Path

# ── Constants ──────────────────────────────────────────────────────────────────

# A ledger data row always starts with MM/DD/YYYY followed by an uppercase code
ENTRY_START = re.compile(r'^(\d{2}/\d{2}/\d{4})\s+([A-Z][A-Z0-9\-]*)\s+')

# Amount is always the last token on the primary line: -$3,697 or $10,000
AMOUNT_END = re.compile(r'(-?\$[\d,]+)\s*$')

# Program name anchor
PROG_RE = re.compile(r'(Certified\s+\S+\s+Welding)')

# Payment period: "481 - 960 hours" style
PP_RE = re.compile(r'^(\d+\s*-\s*\d+\s+hours)\s*')

# Lines to skip entirely (headers, footers, page markers)
SKIP_RES = [re.compile(p, re.IGNORECASE) for p in [
    r'KENTUCKY WELDING INSTITUTE',
    r'^\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}:\d{2}',
    r'^1828 Maysville',
    r'^Phone:',
    r'^https?://',
    r'^Date\s+Code\s+Description',
    r'^Student Ledger\s*-',
    r'^Current Balance',
    r'^Generated:',
    r'^\*\s*\*\s*pending',
    r'Page \d+ of \d+',
]]


def _should_skip(line):
    return any(p.search(line) for p in SKIP_RES)


def _clean_hyphen_break(text):
    return re.sub(r'-\s+', '', text)


def _parse_primary_line(line):
    start_m = ENTRY_START.match(line)
    amt_m = AMOUNT_END.search(line)
    if not start_m or not amt_m:
        return None

    date = start_m.group(1)
    code = start_m.group(2)
    rest_full = line[start_m.end():]
    amount = amt_m.group(1)
    rest = rest_full[:amt_m.start() - start_m.end()].strip()

    prog_m = PROG_RE.search(rest)
    if prog_m:
        description = rest[:prog_m.start()].strip()
        program_base = prog_m.group(1)
        after_prog = rest[prog_m.end():].strip()

        pp_m = PP_RE.match(after_prog)
        if pp_m:
            payment_period = pp_m.group(1).strip()
            notes_start = after_prog[pp_m.end():].strip()
        else:
            parts = after_prog.split(None, 1)
            payment_period = ''
            notes_start = parts[1].strip() if len(parts) > 1 else ''
    else:
        description = rest.strip()
        program_base = ''
        payment_period = ''
        notes_start = ''

    return {
        'date': date,
        'code': code,
        'description': description,
        'program_base': program_base,
        'payment_period': payment_period,
        'notes_start': notes_start,
        'amount': amount,
    }


def _merge_continuation(parsed, continuation_lines):
    notes_parts = [parsed['notes_start']] if parsed['notes_start'] else []
    program = parsed['program_base']

    for raw_line in continuation_lines:
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith('(CPW)'):
            program = (parsed['program_base'] + ' (CPW)') if parsed['program_base'] else '(CPW)'
            after_cpw = line[len('(CPW)'):].strip()
            if after_cpw:
                notes_parts.append(after_cpw)
        else:
            notes_parts.append(line)

    notes_clean = _clean_hyphen_break(' '.join(notes_parts)).strip()

    return {
        'date':           parsed['date'],
        'code':           parsed['code'],
        'description':    parsed['description'],
        'program':        program,
        'payment_period': parsed['payment_period'],
        'notes':          notes_clean,
        'amount':         parsed['amount'],
    }


def extract_ledger_from_lines(lines):
    entries = []
    current = []

    for line in lines:
        if _should_skip(line):
            continue
        if not line.strip():
            continue

        start_m = ENTRY_START.match(line)
        amt_m = AMOUNT_END.search(line)

        if start_m and amt_m:
            if current:
                entries.append(current)
            current = [line]
        elif current:
            current.append(line)

    if current:
        entries.append(current)

    records = []
    for entry_lines in entries:
        parsed = _parse_primary_line(entry_lines[0])
        if parsed is None:
            continue
        records.append(_merge_continuation(parsed, entry_lines[1:]))

    return records


def extract_ledger_folder(source_folder, output_xlsx, on_progress=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    sys.path.insert(0, str(Path(__file__).parent))
    from modules.pdf_extractor import ConfigurableExtractor, DEFAULT_CONFIG

    engine = ConfigurableExtractor(DEFAULT_CONFIG)

    all_files = list(Path(source_folder).glob('*.pdf')) + list(Path(source_folder).glob('*.PDF'))
    seen = set()
    pdfs = []
    for p in sorted(all_files):
        resolved = str(p.resolve()).lower()
        if resolved not in seen:
            seen.add(resolved)
            pdfs.append(p)

    if not pdfs:
        raise RuntimeError(f"No PDF files found in: {source_folder}")

    COLUMNS = ['student_name', 'current_balance', 'source_file',
               'date', 'code', 'description', 'program',
               'payment_period', 'notes', 'amount']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ledger'
    ws.append(COLUMNS)

    header_fill = PatternFill('solid', fgColor='1F4E79')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    total_rows = 0
    files_done = 0

    for pdf_path in pdfs:
        if on_progress:
            on_progress(f'Processing: {pdf_path.name}')
        try:
            lines = engine.read_pdf_lines(str(pdf_path))

            student_name = 'N/A'
            current_balance = 'N/A'
            for line in lines[:12]:
                m = re.search(r'Student Ledger - (.+)', line)
                if m:
                    student_name = m.group(1).strip()
                m2 = re.search(r'Current Balance:\s*(\S+)', line)
                if m2:
                    current_balance = m2.group(1).strip()

            records = extract_ledger_from_lines(lines)
            for rec in records:
                ws.append([
                    student_name, current_balance, pdf_path.name,
                    rec['date'], rec['code'], rec['description'],
                    rec['program'], rec['payment_period'],
                    rec['notes'], rec['amount'],
                ])
                total_rows += 1

            files_done += 1
            if on_progress:
                on_progress(f'  -> {len(records)} rows from {pdf_path.name}')

        except Exception as exc:
            if on_progress:
                on_progress(f'  ERROR {pdf_path.name}: {exc}')

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    ws.freeze_panes = 'A2'
    wb.save(output_xlsx)

    return {'files': files_done, 'rows': total_rows, 'output': output_xlsx}


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: python extract_student_ledger.py <source_folder> <output.xlsx>')
        sys.exit(1)

    result = extract_ledger_folder(
        source_folder=sys.argv[1],
        output_xlsx=sys.argv[2],
        on_progress=print,
    )
    print(f"\nDone! {result['rows']} rows from {result['files']} file(s) -> {result['output']}")
